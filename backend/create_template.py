from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path


# ============================================================
# FILE SETTINGS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_FILE = BASE_DIR / "template.xlsx"


# ============================================================
# EXACT TEMPLATE HEADERS
# ============================================================

HEADERS = [
    "Item name*",
    "Item code",
    "Category",
    "HSN",
    "Default Mrp",
    "Sale price",
    "Online Store Price",
    "Current stock quantity",
    "Minimum stock quantity",
    "Item Location",
    "Tax Rate",
    "Inclusive Of Tax",
    "Base Unit (x)",
    "Secondary Unit (y)",
    "Conversion Rate",
]


# ============================================================
# CREATE TEMPLATE
# ============================================================

def create_template():

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Export Items"


    # ========================================================
    # HEADER
    # ========================================================

    for column_number, header in enumerate(HEADERS, start=1):

        cell = worksheet.cell(
            row=1,
            column=column_number,
            value=header
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="2563EB"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    # ========================================================
    # HEADER HEIGHT
    # ========================================================

    worksheet.row_dimensions[1].height = 28


    # ========================================================
    # COLUMN WIDTHS
    # ========================================================

    column_widths = {
        1: 38,   # Item name
        2: 16,   # Item code
        3: 16,   # Category
        4: 12,   # HSN
        5: 15,   # Default Mrp
        6: 15,   # Sale price
        7: 20,   # Online Store Price
        8: 23,   # Current stock quantity
        9: 23,   # Minimum stock quantity
        10: 18,  # Item Location
        11: 12,  # Tax Rate
        12: 18,  # Inclusive Of Tax
        13: 17,  # Base Unit
        14: 20,  # Secondary Unit
        15: 18,  # Conversion Rate
    }


    for column_number, width in column_widths.items():

        worksheet.column_dimensions[
            get_column_letter(column_number)
        ].width = width


    # ========================================================
    # BORDER
    # ========================================================

    thin_border = Border(
        left=Side(style="thin", color="D9E1EA"),
        right=Side(style="thin", color="D9E1EA"),
        top=Side(style="thin", color="D9E1EA"),
        bottom=Side(style="thin", color="D9E1EA"),
    )


    # ========================================================
    # FORMAT EMPTY DATA ROWS
    # ========================================================

    for row in range(2, 102):

        for column in range(1, 16):

            cell = worksheet.cell(
                row=row,
                column=column
            )

            cell.border = thin_border

            cell.alignment = Alignment(
                vertical="center"
            )


    # ========================================================
    # NUMBER FORMATS
    # ========================================================

    for row in range(2, 102):

        # MRP
        worksheet.cell(row, 5).number_format = "0.00"

        # Sale Price
        worksheet.cell(row, 6).number_format = "0.00"

        # Online Store Price
        worksheet.cell(row, 7).number_format = "0.00"

        # Current Stock
        worksheet.cell(row, 8).number_format = "0"

        # Minimum Stock
        worksheet.cell(row, 9).number_format = "0"

        # Tax Rate
        worksheet.cell(row, 11).number_format = "0.00"

        # Conversion Rate
        worksheet.cell(row, 15).number_format = "0.00"


    # ========================================================
    # FREEZE HEADER
    # ========================================================

    worksheet.freeze_panes = "A2"


    # ========================================================
    # FILTER
    # ========================================================

    worksheet.auto_filter.ref = "A1:O101"


    # ========================================================
    # SAVE
    # ========================================================

    workbook.save(OUTPUT_FILE)

    print("=" * 60)
    print("TEMPLATE CREATED SUCCESSFULLY")
    print("=" * 60)
    print(f"File: {OUTPUT_FILE}")
    print(f"Columns: {len(HEADERS)}")
    print("=" * 60)


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    create_template()