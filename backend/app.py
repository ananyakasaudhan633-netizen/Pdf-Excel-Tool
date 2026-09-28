from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
from pypdf import PdfReader
from openpyxl import load_workbook

import os
import re
import tempfile


# ============================================================
# FLASK APP
# ============================================================

app = Flask(__name__)


# ============================================================
# CORS
# ============================================================

CORS(
    app,
    resources={
        r"/api/*": {
            "origins": [
                "http://localhost:5173",
                "http://localhost:5174",
                "http://127.0.0.1:5173",
                "http://127.0.0.1:5174",
            ]
        }
    },
)


# ============================================================
# FILE SETTINGS
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

TEMPLATE_PATH = os.path.join(
    BASE_DIR,
    "template.xlsx"
)


# ============================================================
# NUMBER PATTERN
# ============================================================

NUMBER_PATTERN = r"\d+(?:,\d{3})*(?:\.\d+)?"


# ============================================================
# CLEAN NUMBER
# ============================================================

def clean_number(value):
    """Convert text number into float."""

    if value is None:
        return 0.0

    value = str(value).strip()
    value = value.replace(",", "")
    value = value.replace("₹", "")
    value = value.replace("%", "")

    try:
        return float(value)
    except (ValueError, TypeError):
        return 0.0


# ============================================================
# EXTRACT PDF TEXT
# ============================================================

def extract_pdf_text(pdf_path):
    """Extract all text from PDF."""

    reader = PdfReader(pdf_path)

    pages = []

    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)

    return "\n".join(pages)


# ============================================================
# EXTRACT PRODUCTS
# ============================================================

def extract_products(text):

    products = []

    excluded_words = [
        "PACKING",
        "TRANSPORT",
        "TRANSPORTATION",
        "FREIGHT",
        "ROUND OFF",
        "ROUNDOFF",
        "CGST",
        "SGST",
        "IGST",
        "GRAND TOTAL",
        "SUB TOTAL",
        "SUBTOTAL",
    ]

    for raw_line in text.splitlines():

        line = " ".join(
            raw_line.strip().split()
        )

        if not line:
            continue

        upper_line = line.upper()

        if any(
            word in upper_line
            for word in excluded_words
        ):
            continue

        # Product row starts like:
        # 1. Product Name

        serial_match = re.match(
            r"^(\d+)\.\s+(.+)$",
            line
        )

        if not serial_match:
            continue

        serial_no = int(
            serial_match.group(1)
        )

        remaining = serial_match.group(2).strip()

        # ----------------------------------------------------
        # Quantity + Unit
        # ----------------------------------------------------

        qty_match = re.search(
            rf"({NUMBER_PATTERN})\s+([A-Za-z][A-Za-z.]*)",
            remaining
        )

        if not qty_match:
            continue

        qty = clean_number(
            qty_match.group(1)
        )

        unit = qty_match.group(2).strip()

        if qty <= 0:
            continue

        product_name = remaining[
            :qty_match.start()
        ].strip()

        if not product_name:
            continue

        # ----------------------------------------------------
        # Numeric values after quantity and unit
        # ----------------------------------------------------

        numeric_text = remaining[
            qty_match.end():
        ]

        numeric_values = [
            clean_number(value)
            for value in re.findall(
                NUMBER_PATTERN,
                numeric_text
            )
        ]

        if len(numeric_values) < 3:
            continue

        # ----------------------------------------------------
        # Invoice structure
        #
        # First numeric = MRP
        # Second-last = Unit Price
        # Last = Amount
        # ----------------------------------------------------

        list_price = numeric_values[0]
        price = numeric_values[-2]
        amount = numeric_values[-1]

        if list_price <= 0:
            continue

        if price <= 0:
            continue

        if amount <= 0:
            continue

        products.append(
            {
                "serial_no": serial_no,
                "name": product_name,
                "qty": qty,
                "unit": unit,
                "list_price": list_price,
                "price": price,
                "amount": amount,
            }
        )

    return products


# ============================================================
# EXTRACT GRAND TOTAL
# ============================================================

def extract_grand_total(text):

    for raw_line in text.splitlines():

        line = " ".join(
            raw_line.strip().split()
        )

        if "GRAND TOTAL" not in line.upper():
            continue

        values = [
            clean_number(value)
            for value in re.findall(
                NUMBER_PATTERN,
                line
            )
        ]

        values = [
            value
            for value in values
            if value > 0
        ]

        if len(values) >= 2:
            return max(values)

        if len(values) == 1:
            return values[0]

    # Fallback
    products = extract_products(text)

    if products:
        return round(
            sum(
                product["amount"]
                for product in products
            ),
            2
        )

    return 0.0


# ============================================================
# EXTRACT TOTAL UNITS
# ============================================================

def extract_total_units(text):

    for raw_line in text.splitlines():

        line = " ".join(
            raw_line.strip().split()
        )

        if "GRAND TOTAL" not in line.upper():
            continue

        values = [
            clean_number(value)
            for value in re.findall(
                NUMBER_PATTERN,
                line
            )
        ]

        values = [
            value
            for value in values
            if value > 0
        ]

        if len(values) >= 2:
            return min(values)

        if len(values) == 1:
            return values[0]

    return 0.0


# ============================================================
# EXTRA CHARGE PER PCS
# ============================================================

def calculate_extra_charge_per_unit(
    total_extra_charge,
    unit_rate,
    total_amount
):
    """
    Formula:

    Extra Charge / PCS =
    (TOTAL EXTRA CHARGE * 100 / TOTAL AMOUNT)
    * UNIT RATE / 100

    Equivalent to:

    TOTAL EXTRA CHARGE * UNIT RATE / TOTAL AMOUNT
    """

    if total_amount <= 0:
        return 0.0

    extra_charge_per_unit = (
        (
            total_extra_charge * 100
        )
        / total_amount
    ) * (
        unit_rate / 100
    )

    return round(
        extra_charge_per_unit,
        3
    )


# ============================================================
# SELLING PRICE
# ============================================================

def calculate_selling_price(
    unit_price,
    extra_charge_per_unit,
    profit_percentage
):
    """
    Cost Price =
    Unit Price + Extra Charge / PCS

    Profit =
    Cost Price * Profit % / 100

    Selling Price =
    Cost Price + Profit
    """

    cost_price = (
        unit_price
        + extra_charge_per_unit
    )

    profit_amount = (
        cost_price
        * profit_percentage
        / 100
    )

    selling_price = (
        cost_price
        + profit_amount
    )

    return round(
        selling_price,
        2
    )


# ============================================================
# CREATE EXCEL
# ============================================================

def create_excel(
    products,
    output_path,
    invoice_total,
    item_code,
    category,
    hsn,
    item_location,
    minimum_stock,
    transport_charge,
    gst_amount,
    any_extra_charge,
    profit_percentage,
    gst_rate
):

    if not os.path.exists(
        TEMPLATE_PATH
    ):
        raise FileNotFoundError(
            "template.xlsx not found at: "
            + TEMPLATE_PATH
        )

    # --------------------------------------------------------
    # Load existing template
    # --------------------------------------------------------

    workbook = load_workbook(
        TEMPLATE_PATH
    )

    if "Export Items" in workbook.sheetnames:
        worksheet = workbook["Export Items"]
    else:
        worksheet = workbook.active

    # --------------------------------------------------------
    # Remove old data but keep header
    # --------------------------------------------------------

    if worksheet.max_row >= 2:

        worksheet.delete_rows(
            2,
            worksheet.max_row - 1
        )

    # --------------------------------------------------------
    # Total Extra Charge
    # --------------------------------------------------------

    total_extra_charge = (
        transport_charge
        + gst_amount
        + any_extra_charge
    )

    # --------------------------------------------------------
    # GST text
    # --------------------------------------------------------

    if gst_rate > 0:
        tax_rate_text = f"GST@{gst_rate:g}%"
    else:
        tax_rate_text = "GST@0%"

    # --------------------------------------------------------
    # Add products
    # --------------------------------------------------------

    for product in products:

        # ----------------------------------------------------
        # IMPORTANT:
        # Formula uses UNIT RATE.
        #
        # Here Unit Rate = PDF Unit Price
        # ----------------------------------------------------

        unit_rate = product["price"]

        extra_charge_per_unit = (
            calculate_extra_charge_per_unit(
                total_extra_charge,
                unit_rate,
                invoice_total
            )
        )

        selling_price = (
            calculate_selling_price(
                unit_rate,
                extra_charge_per_unit,
                profit_percentage
            )
        )

        # Default MRP = Selling Price x 2

        default_mrp = round(
            selling_price * 2,
            2
        )

        row_number = (
            worksheet.max_row + 1
        )

        # ----------------------------------------------------
        # EXACT TEMPLATE A-O MAPPING
        # ----------------------------------------------------

        values = {

            # A - Item name*
            1: product["name"],

            # B - Item code
            2: item_code,

            # C - Category
            3: category,

            # D - HSN
            4: hsn,

            # E - Default MRP
            5: default_mrp,

            # F - Sale price
            6: selling_price,

            # G - Purchase price
            7: unit_rate,

            # H - Opening stock quantity
            8: product["qty"],

            # I - Minimum stock quantity
            9: minimum_stock,

            # J - Item Location
            10: item_location,

            # K - Tax Rate
            11: tax_rate_text,

            # L - Inclusive Of Tax
            12: "Inclusive",

            # M - Base Unit (x)
            13: product["unit"],

            # N - Secondary Unit (y)
            14: "",

            # O - Conversion Rate (n)
            15: "",
        }

        for column_number, value in values.items():

            worksheet.cell(
                row=row_number,
                column=column_number,
                value=value
            )

    # --------------------------------------------------------
    # Save Excel
    # --------------------------------------------------------

    workbook.save(
        output_path
    )


# ============================================================
# HOME
# ============================================================

@app.route("/")
def home():

    return jsonify(
        {
            "success": True,
            "message":
                "PDF to Excel Import Tool API is running"
        }
    )


# ============================================================
# HEALTH CHECK
# ============================================================

@app.route(
    "/api/health",
    methods=["GET"]
)
def health():

    return jsonify(
        {
            "success": True,
            "status": "ok",
            "message":
                "PDF to Excel backend is running"
        }
    )


# ============================================================
# GENERATE EXCEL
# ============================================================

@app.route(
    "/api/generate-excel",
    methods=["POST"]
)
def generate_excel():

    pdf_file = request.files.get("pdf")

    if not pdf_file:

        return jsonify(
            {
                "success": False,
                "error":
                    "Please upload a PDF file."
            }
        ), 400

    if not pdf_file.filename:

        return jsonify(
            {
                "success": False,
                "error":
                    "Invalid PDF filename."
            }
        ), 400

    # ========================================================
    # USER INPUTS
    # ========================================================

    item_code = request.form.get(
        "item_code",
        ""
    ).strip()

    category = request.form.get(
        "category",
        ""
    ).strip()

    hsn = request.form.get(
        "hsn",
        ""
    ).strip()

    item_location = request.form.get(
        "item_location",
        ""
    ).strip()

    minimum_stock = clean_number(
        request.form.get(
            "minimum_stock",
            3
        )
    )

    transport_charge = clean_number(
        request.form.get(
            "transport_charge",
            0
        )
    )

    gst_amount = clean_number(
        request.form.get(
            "gst_amount",
            0
        )
    )

    any_extra_charge = clean_number(
        request.form.get(
            "any_extra_charge",
            0
        )
    )

    profit_percentage = clean_number(
        request.form.get(
            "profit_percentage",
            0
        )
    )

    gst_rate = clean_number(
        request.form.get(
            "gst_rate",
            0
        )
    )

    # ========================================================
    # BASIC VALIDATION
    # ========================================================

    if not item_code:

        return jsonify(
            {
                "success": False,
                "error":
                    "Item Code is required."
            }
        ), 400

    if not category:

        return jsonify(
            {
                "success": False,
                "error":
                    "Category is required."
            }
        ), 400

    if not hsn:

        return jsonify(
            {
                "success": False,
                "error":
                    "HSN is required."
            }
        ), 400

    if transport_charge < 0:

        return jsonify(
            {
                "success": False,
                "error":
                    "Transport Charge cannot be negative."
            }
        ), 400

    if gst_amount < 0:

        return jsonify(
            {
                "success": False,
                "error":
                    "GST Amount cannot be negative."
            }
        ), 400

    if any_extra_charge < 0:

        return jsonify(
            {
                "success": False,
                "error":
                    "Any Extra Charge cannot be negative."
            }
        ), 400

    if profit_percentage < 0:

        return jsonify(
            {
                "success": False,
                "error":
                    "Profit Percentage cannot be negative."
            }
        ), 400

    temp_pdf = None

    try:

        # ====================================================
        # SAVE TEMPORARY PDF
        # ====================================================

        with tempfile.NamedTemporaryFile(
            delete=False,
            suffix=".pdf"
        ) as temp_file:

            pdf_file.save(
                temp_file.name
            )

            temp_pdf = temp_file.name

        # ====================================================
        # EXTRACT PDF TEXT
        # ====================================================

        text = extract_pdf_text(
            temp_pdf
        )

        # ====================================================
        # EXTRACT PRODUCTS
        # ====================================================

        products = extract_products(
            text
        )

        if not products:

            return jsonify(
                {
                    "success": False,
                    "error":
                        "No product rows could be extracted from the PDF."
                }
            ), 400

        # ====================================================
        # EXTRACT INVOICE TOTAL
        # ====================================================

        total_amount = extract_grand_total(
            text
        )

        if total_amount <= 0:

            return jsonify(
                {
                    "success": False,
                    "error":
                        "Invoice total could not be detected from the PDF."
                }
            ), 400

        # ====================================================
        # TOTAL UNITS
        # ====================================================

        displayed_total_units = (
            extract_total_units(text)
        )

        extracted_total_units = sum(
            product["qty"]
            for product in products
        )

        # ====================================================
        # TOTAL EXTRA CHARGE
        # ====================================================

        total_extra_charge = (
            transport_charge
            + gst_amount
            + any_extra_charge
        )

        # ====================================================
        # PRINT SUMMARY
        # ====================================================

        print()
        print("= - app.py:885" * 60)
        print("PDF PROCESSING SUMMARY - app.py:886")
        print("= - app.py:887" * 60)

        print(
            f"Products extracted: {len(products)}"
        )

        print(
            f"Invoice total: {total_amount}"
        )

        print(
            f"PDF displayed units: "
            f"{displayed_total_units}"
        )

        print(
            f"Extracted product units: "
            f"{extracted_total_units}"
        )

        print(
            f"Unit difference: "
            f"{displayed_total_units - extracted_total_units}"
        )

        print(
            f"Transport charge: "
            f"{transport_charge}"
        )

        print(
            f"GST amount: "
            f"{gst_amount}"
        )

        print(
            f"Any extra charge: "
            f"{any_extra_charge}"
        )

        print(
            f"TOTAL EXTRA CHARGE: "
            f"{total_extra_charge}"
        )

        print(
            f"Profit percentage: "
            f"{profit_percentage}%"
        )

        print(
            f"GST rate: "
            f"{gst_rate}%"
        )

        print("= - app.py:942" * 60)

        # ====================================================
        # FIRST PRODUCT CALCULATION CHECK
        # ====================================================

        first_product = products[0]

        first_unit_rate = (
            first_product["price"]
        )

        first_extra_per_unit = (
            calculate_extra_charge_per_unit(
                total_extra_charge,
                first_unit_rate,
                total_amount
            )
        )

        first_selling_price = (
            calculate_selling_price(
                first_unit_rate,
                first_extra_per_unit,
                profit_percentage
            )
        )

        first_default_mrp = round(
            first_selling_price * 2,
            2
        )

        print()
        print("FIRST PRODUCT CHECK - app.py:976")
        print("= - app.py:977" * 60)

        print(
            f"Product: "
            f"{first_product['name']}"
        )

        print(
            f"Unit Price: "
            f"{first_unit_rate}"
        )

        print(
            f"Extra Charge/PCS: "
            f"{first_extra_per_unit}"
        )

        print(
            f"Selling Price: "
            f"{first_selling_price}"
        )

        print(
            f"Default MRP: "
            f"{first_default_mrp}"
        )

        print("= - app.py:1004" * 60)

        # ====================================================
        # CREATE EXCEL
        # ====================================================

        output_path = os.path.join(
            BASE_DIR,
            "product_import.xlsx"
        )

        create_excel(
            products=products,
            output_path=output_path,
            invoice_total=total_amount,
            item_code=item_code,
            category=category,
            hsn=hsn,
            item_location=item_location,
            minimum_stock=minimum_stock,
            transport_charge=transport_charge,
            gst_amount=gst_amount,
            any_extra_charge=any_extra_charge,
            profit_percentage=profit_percentage,
            gst_rate=gst_rate
        )

        # ====================================================
        # SEND EXCEL
        # ====================================================

        return send_file(
            output_path,
            as_attachment=True,
            download_name="product_import.xlsx",
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

    except Exception as exc:

        print(
            "ERROR:",
            str(exc)
        )

        return jsonify(
            {
                "success": False,
                "error": str(exc)
            }
        ), 500

    finally:

        if (
            temp_pdf
            and os.path.exists(temp_pdf)
        ):

            try:
                os.remove(temp_pdf)

            except OSError:
                pass


# ============================================================
# RUN SERVER
# ============================================================

if __name__ == "__main__":

    print()
    print("= - app.py:1080" * 60)
    print("PDF TO EXCEL BACKEND - app.py:1081")
    print("= - app.py:1082" * 60)

    print(
        "Template:",
        TEMPLATE_PATH
    )

    print(
        "Server:",
        "http://127.0.0.1:5000"
    )

    print("= - app.py:1094" * 60)

    app.run(
        host="127.0.0.1",
        port=5000,
        debug=True
    )